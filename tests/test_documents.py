from docx import Document

from backend.document.export_service import build_docx_bytes, build_pdf_bytes, build_txt
from backend.models import GeneratedDocument


def sample() -> GeneratedDocument:
    return GeneratedDocument(
        title="Test Agreement",
        introduction="A short test document.",
        sections=["PARTIES\nAlice and Bob.", "TERM\nThe term is one year."],
        clauses=["Payment within 30 days", "Confidentiality applies"],
        closing="Review before signing.",
        warnings=[],
    )


def test_txt_export() -> None:
    data = build_txt(sample())
    assert b"Test Agreement" in data


def test_docx_export() -> None:
    data = build_docx_bytes(sample())
    assert data.startswith(b"PK")
    doc = Document(__import__("io").BytesIO(data))
    assert any("Test Agreement" in p.text for p in doc.paragraphs)


def test_pdf_export() -> None:
    data = build_pdf_bytes(sample())
    assert data.startswith(b"%PDF")
