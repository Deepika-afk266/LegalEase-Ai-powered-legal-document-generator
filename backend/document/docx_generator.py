from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from backend.models import GeneratedDocument
from backend.utils.text_utils import sanitize_text


def _set_cell_shading(cell, fill: str = "E8EEF7") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def build_docx(document_data: GeneratedDocument, logo_bytes: bytes | None = None) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)
    styles["Heading 1"].font.name = "Times New Roman"
    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 1"].font.bold = True

    if logo_bytes:
        try:
            doc.add_picture(BytesIO(logo_bytes), width=Inches(1.2))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            pass

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(sanitize_text(document_data.title))
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    if document_data.introduction:
        p = doc.add_paragraph(sanitize_text(document_data.introduction))
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for section_text in document_data.sections:
        parts = section_text.split("\n", 1)
        heading = parts[0].strip()
        body = parts[1].strip() if len(parts) > 1 else ""
        h = doc.add_paragraph()
        h.style = "Heading 1"
        h.add_run(sanitize_text(heading))
        if body:
            p = doc.add_paragraph(sanitize_text(body))
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if document_data.clauses:
        h = doc.add_paragraph()
        h.style = "Heading 1"
        h.add_run("KEY TERMS")
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        header = table.rows[0].cells
        header[0].text = "#"
        header[1].text = "Term / Condition"
        _set_cell_shading(header[0])
        _set_cell_shading(header[1])
        for idx, clause in enumerate(document_data.clauses, 1):
            row = table.add_row().cells
            row[0].text = str(idx)
            row[1].text = sanitize_text(clause)

    if document_data.closing:
        p = doc.add_paragraph(sanitize_text(document_data.closing))
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Generated with LegalEase • Draft for review")

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
